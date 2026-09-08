"""Append a test record to the data-search test Excel tracker."""
import argparse, os, sys
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

EXCEL_PATH = r"C:\Users\ziyao.yang.int\Desktop\yzy-0622\数据搜索\robowork测试记录.xlsx"

HEADERS = ["问题", "查询耗时", "查询结果", "结果是否正确", "调用API", "API数据是否支持", "备注"]
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
DATA_FONT = Font(name="微软雅黑", size=10)
DATA_ALIGN = Alignment(vertical="center", wrap_text=True)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def main():
    parser = argparse.ArgumentParser(description="Append a data-search test record to Excel")
    parser.add_argument("--question", required=True, help="The original question asked")
    parser.add_argument("--time", default="", help="Query elapsed time, e.g. 2m48s")
    parser.add_argument("--result", default="", help="Result placeholder (leave empty)")
    parser.add_argument("--correct", default="待验证", help="Correct / 未查询到 / 待验证")
    parser.add_argument("--api", default="", help="APIs called, semicolon-separated")
    parser.add_argument("--supported", default="", help="支持 / 不支持 / 部分支持")
    parser.add_argument("--notes", default="", help="Remarks (【缺什么】【现状】【建议】 required when supported != 支持)")
    args = parser.parse_args()

    # 🔴 强制校验：部分支持 / 不支持时，备注必填且必须含三要素
    # 把命令行传入的字面 \n 转为真正换行，使 Excel 单元格内三要素分行显示
    args.notes = args.notes.replace("\\n", "\n")

    if args.supported in ("部分支持", "不支持"):
        missing_tags = []
        for tag in ("【缺什么】", "【现状】", "【建议】"):
            if tag not in args.notes:
                missing_tags.append(tag)
        if missing_tags:
            print(
                f"[ERROR] supported={args.supported} 时备注为必填，且须含【缺什么】【现状】【建议】三要素。"
                f" 缺失标签: {', '.join(missing_tags)}",
                file=sys.stderr,
            )
            sys.exit(1)

    if os.path.exists(EXCEL_PATH):
        wb = load_workbook(EXCEL_PATH)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "测试用例"
        # Write headers
        for col_idx, h in enumerate(HEADERS, 1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGN
            cell.border = THIN_BORDER
        # Column widths
        ws.column_dimensions["A"].width = 50
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 20
        ws.column_dimensions["D"].width = 16
        ws.column_dimensions["E"].width = 35
        ws.column_dimensions["F"].width = 16
        ws.column_dimensions["G"].width = 45

    # Find next empty row
    next_row = ws.max_row + 1

    values = [
        args.question,
        args.time,
        args.result,
        args.correct,
        args.api,
        args.supported,
        args.notes,
    ]
    for col_idx, val in enumerate(values, 1):
        cell = ws.cell(row=next_row, column=col_idx, value=val)
        cell.font = DATA_FONT
        cell.alignment = DATA_ALIGN
        cell.border = THIN_BORDER

    wb.save(EXCEL_PATH)
    print(f"[OK] Written to {EXCEL_PATH}, row {next_row}")
    print(f"  question={args.question}")
    print(f"  time={args.time}  correct={args.correct}  api={args.api}  supported={args.supported}")
    if args.notes:
        print(f"  notes={args.notes}")


if __name__ == "__main__":
    main()
